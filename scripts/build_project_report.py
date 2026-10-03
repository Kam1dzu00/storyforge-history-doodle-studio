from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "ОТЧЕТ_HISTORY_DOODLE_STUDIO.docx"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 28), ("Heading 1", 18), ("Heading 2", 13)):
    s = styles[name]
    s.font.name = "Arial"
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
    s.paragraph_format.space_after = Pt(8)

# Some Word-compatible renderers add a theme border to the built-in Title
# style. Remove it explicitly so the cover relies on typography and spacing.
title_ppr = styles["Title"].element.get_or_add_pPr()
old_borders = title_ppr.find(qn("w:pBdr"))
if old_borders is not None:
    title_ppr.remove(old_borders)

def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)

def borders(table):
    tbl_pr = table._tbl.tblPr
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        node.append(e)
    tbl_pr.append(node)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, "25354A")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        if widths: c.width = Inches(widths[i])
    header_props = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_props.append(repeat)
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        row_props = t.rows[-1]._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if n % 2: shade(cells[i], "F2F5F8")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.add_run(str(value))
            if widths: cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p

if (ROOT / "assets/app_icon.png").exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / "assets/app_icon.png"), width=Inches(1.15))

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Отчет по проекту History Doodle Studio")
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("Локальная платформа создания исторических роликов на Apple Silicon")
run.bold = True
run.font.size = Pt(13)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Версия 1.0   |   Состояние на 3 октября 2026 года   |   macOS arm64").font.color.rgb = RGBColor(90, 100, 112)

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Главный результат. ").bold = True
p.add_run("Собрано автономное настольное приложение, которое принимает тему или готовый английский сценарий, создаёт раскадровку, генерирует иллюстрации, записывает голос рассказчика, синхронизирует субтитры и собирает готовый ролик MP4. Основной рабочий цикл выполняется локально на Mac без облачной генерации.")

doc.add_heading("1 Назначение проекта", level=1)
doc.add_paragraph("History Doodle Studio предназначен для серийного выпуска англоязычных исторических объясняющих роликов. Визуальная концепция закреплена как взрослый редакционный исторический мультфильм: человеческие пропорции, простые читаемые лица, тёмный рисованный контур, приглушённая земная палитра и текстура состаренной бумаги. Детская книжная эстетика, чиби-пропорции, коллажи и текст внутри сгенерированных сцен исключаются промптами и проверками рабочего процесса.")

doc.add_heading("2 Что установлено", level=1)
table(
    ["Компонент", "Назначение", "Настройка"],
    [
        ["Qwen3 4B Instruct 4 bit", "Сценарист и режиссёр раскадровки", "MLX, локальная модель"],
        ["FLUX 2 Klein 4B", "Генерация кадров 16 к 9", "MFLUX, 4 шага, 1024 на 576"],
        ["Kokoro 82M", "Английский голос рассказчика", "Локальный TTS, скорость 1.03"],
        ["Whisper Base MLX", "Тайминг слов для субтитров", "Локальная транскрипция"],
        ["FFmpeg", "Монтаж, звук и финальный MP4", "H.264, AAC, 1920 на 1080, 30 fps"],
        ["PySide6", "Русский графический интерфейс", "Настольное приложение macOS"],
        ["Python 3.12", "Среда исполнения", "Изолированное окружение проекта"],
    ], [1.65, 2.85, 1.95]
)
doc.add_paragraph("Фактический общий объём локального кэша моделей составляет около 25 ГБ. Модели хранятся внутри проекта, а не разбрасываются по системным папкам пользователя.")

doc.add_heading("3 Как работает полный цикл", level=1)
steps = [
    ("1", "Ввод", "Пользователь вводит тему либо вставляет готовый английский текст рассказчика."),
    ("2", "Сценарий", "Qwen создаёт сценарий только для режима идеи. Готовый сценарий не переписывается."),
    ("3", "Раскадровка", "Текст делится на разделы, сцены и отдельные кадры длительностью от 3 до 10 секунд."),
    ("4", "Озвучка", "Kokoro создаёт английскую дорожку. Паузы очищаются и громкость нормализуется."),
    ("5", "Тайминг", "Whisper определяет начало и конец каждого слова для точных субтитров."),
    ("6", "Иллюстрации", "FLUX создаёт одно цельное изображение для каждого кадра с фиксированным стилем и seed."),
    ("7", "Монтаж", "Кадры сводятся с мягкими переходами, спокойным движением, озвучкой и субтитрами."),
    ("8", "Экспорт", "Готовый ролик сохраняется в MP4 вместе с версией без вшитых субтитров."),
]
table(["Этап", "Название", "Результат"], steps, [0.55, 1.35, 4.55])

doc.add_heading("4 Настройки генерации", level=1)
table(
    ["Параметр", "Текущее значение"],
    [
        ["Формат кадра", "16 к 9"],
        ["Размер генерации", "1024 на 576 пикселей"],
        ["Финальное видео", "1920 на 1080, 30 кадров в секунду"],
        ["Длительность кадра", "3–10 секунд, обычно 4–6 секунд"],
        ["Движение", "Статичный кадр или плавное изменение масштаба до 1.5 процента"],
        ["Переход", "Мягкое растворение 0.32 секунды"],
        ["Субтитры", "Английские, синхронизированные по словам, вшитые в видео"],
        ["Звук", "AAC 192 кбит в секунду, голос нормализован примерно до минус 16 LUFS"],
    ], [2.1, 4.35]
)

doc.add_heading("5 Защита визуального стиля", level=1)
doc.add_paragraph("Каждый запрос изображения начинается с неизменяемого описания стиля. К нему добавляются герой, локация, действие, композиция и исторические предметы конкретного кадра. Отдельный отрицательный промпт запрещает фотографичность, 3D, детскую стилизацию, чиби, несколько панелей, коллаж, логотипы, буквы и цифры. Случайное число seed сохраняется для воспроизводимости.")
bullet("Кнопка Повторить точно использует тот же seed и описание.")
bullet("Кнопка Новый вариант изменяет seed, но сохраняет сцену и стиль.")
bullet("Описание выбранного кадра можно поправить до повторной генерации.")
bullet("Файлы уже созданных кадров переиспользуются, если их промпт не изменился.")

doc.add_heading("6 Интерфейс приложения", level=1)
doc.add_paragraph("Интерфейс полностью переведён на русский язык. Слева находятся проекты и состояние этапов, в центре — ввод материала, длительность, голос и главный запуск, справа — раскадровка, просмотр изображения и ручное управление кадром. Для обычной работы достаточно заполнить материал и нажать кнопку Создать готовое видео весь процесс.")
table(
    ["Элемент", "Что делает"],
    [
        ["Новый ролик", "Создаёт отдельную папку проекта с собственными файлами и состоянием."],
        ["У меня есть идея", "Передаёт тему локальному сценаристу."],
        ["У меня есть сценарий", "Использует вставленный текст без переписывания."],
        ["Создать готовое видео", "Последовательно запускает все незавершённые этапы."],
        ["Ручное управление", "Запускает только выбранный этап для проверки или исправления."],
        ["Пересобрать раскадровку", "Повторно анализирует изменённый сценарий и помечает зависимые материалы устаревшими."],
        ["Проверить систему", "Проверяет модели, Apple Silicon и FFmpeg."],
    ], [2.1, 4.35]
)

doc.add_heading("7 Где находятся файлы", level=1)
table(
    ["Содержимое", "Путь"],
    [
        ["Главная папка", str(ROOT)],
        ["Приложение", str(ROOT / "HistoryDoodle Studio.app")],
        ["Модели", str(ROOT / "models")],
        ["Проекты", str(ROOT / "projects")],
        ["Стиль", str(ROOT / "assets/style")],
        ["Журнал ошибок", str(ROOT / "logs/app.log")],
    ], [1.55, 4.9]
)
doc.add_paragraph("В каждом проекте используются папки planning, audio, images, subtitles и render. Итоговый ролик называется final_video.mp4. Версия без вшитых субтитров называется final_video_no_subtitles.mp4.")

doc.add_heading("8 Как запускать и создавать серию", level=1)
for text in [
    "Дважды нажмите HistoryDoodle Studio.app.",
    "Нажмите Новый ролик.",
    "Введите название и выберите вкладку идеи либо готового сценария.",
    "Укажите продолжительность и голос рассказчика.",
    "Нажмите Создать готовое видео весь процесс.",
    "Не закрывайте приложение и не переводите Mac в сон во время генерации.",
    "После завершения нажмите Открыть готовое видео.",
]:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)

doc.add_heading("9 Что было исправлено", level=1)
bullet("Заменена тяжёлая локальная видеогенерация на надёжный конвейер качественных изображений, голоса и монтажа.")
bullet("Добавлена строгая JSON схема, проверка ответа режиссёра и две попытки автоматического исправления структуры.")
bullet("Исправлены порядок сцен, соответствие оригинальному тексту и автоматическая продолжительность кадров.")
bullet("Убраны коллажи, панели и текстовые элементы внутри изображений.")
bullet("Убраны сильная тряска и боковые панорамы. Масштабирование ограничено 1.5 процента.")
bullet("Добавлены плавные переходы 0.32 секунды без нарушения синхронизации голоса.")
bullet("Исправлены ошибки счётчика кадров FFmpeg и абсолютных путей при сборке клипов.")
bullet("Добавлено возобновление этапов и повторное использование готовых результатов.")
bullet("Создан русский интерфейс и полноценный ярлык macOS с собственной иконкой.")

doc.add_heading("10 Проверенный тестовый проект", level=1)
table(
    ["Проверка", "Результат"],
    [
        ["Тема", "What Did Vikings Actually Do All Day"],
        ["Продолжительность", "47.461 секунды"],
        ["Иллюстрации", "10 готовых кадров"],
        ["Видео", "H.264, 1920 на 1080, 30 fps"],
        ["Аудио", "AAC, 48 кГц, моно"],
        ["Субтитры", "42 синхронизированных события"],
        ["Автоматические тесты", "15 тестов пройдено"],
        ["Полный повторный запуск", "Успешно, готовые этапы не пересчитываются"],
    ], [2.1, 4.35]
)

doc.add_heading("11 Ограничения и ожидаемое время", level=1)
doc.add_paragraph("MacBook Pro M1 Pro с 16 ГБ памяти способен выполнять весь процесс, но генерация иллюстраций идёт последовательно. Полная серия занимает существенно больше времени, чем тестовый фрагмент. Скорость зависит от количества кадров, температуры компьютера и длины сценария. Во время работы лучше подключить питание, обеспечить охлаждение и отключить переход в сон.")
doc.add_paragraph("Модель 4B может иногда изменить мелкую деталь одежды, предмет или лицо второстепенного персонажа. Для этого предусмотрен просмотр кадров и локальная перегенерация. Генератор не заменяет финальную редакторскую проверку исторических фактов и визуальной целостности.")

doc.add_heading("12 Обслуживание", level=1)
bullet("Не перемещайте отдельно папки models, assets, src или .venv: приложение использует их как единый комплект.")
bullet("Готовые проекты можно архивировать копированием всей папки из projects.")
bullet("При нехватке диска сначала удаляйте ненужные папки проектов, а не модели рабочего комплекта.")
bullet("При ошибке откройте Проверить систему и затем передайте файл logs/app.log для диагностики.")
bullet("Исходный код сохранён в Git. Стабильная версия монтажа зафиксирована отдельным коммитом.")

doc.add_heading("13 Итог", level=1)
doc.add_paragraph("History Doodle Studio готов к самостоятельному созданию серий. Пользовательский путь сведён к вводу идеи или сценария и одному запуску. Все основные компоненты локальные, проекты сохраняются между запусками, отдельные этапы можно повторять, а итоговый файл соответствует стандартному формату для публикации на видеоплатформах.")

for s in doc.sections:
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("History Doodle Studio   •   Технический отчёт")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(120, 125, 132)

doc.core_properties.title = "Отчет по проекту History Doodle Studio"
doc.core_properties.subject = "Архитектура, модели, настройка и рабочий процесс"
doc.core_properties.author = "History Doodle Studio"
doc.save(OUT)
print(OUT)
